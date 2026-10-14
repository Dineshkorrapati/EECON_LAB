"""
Compare the manual coding of the validation sample with the model's labels
and compute precision, recall and F1 for every label.

Input:  validation/manual_coding_completed.xlsx, validation/model_labels_key.xlsx
Output: results/tables/validation_results.xlsx
Run:    python code/validation/03_evaluate_manual_coding.py
"""

import re
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
KEY_FILE = ROOT / "validation" / "model_labels_key.xlsx"
CODER_FILES = {"Manual": ROOT / "validation" / "manual_coding_completed.xlsx"}
OUTPUT_FILE = ROOT / "results" / "tables" / "validation_results.xlsx"

TOPICS = ["nuclear", "energy_prices", "energy_security", "climate_policy", "renewables_and_grid", "coal_and_regions"]
RISKS = ["environmental_risk", "geopolitical_risk", "social_risk"]
SOURCES = ["wind", "natural_gas", "nuclear", "solar", "renewables_general", "coal"]
LABELS = [("Topic", l) for l in TOPICS] + [("Risk", l) for l in RISKS] + [("Source", l) for l in SOURCES]


def read_coder(path):
    ws = load_workbook(path, data_only=True)["Coding"]
    rows = {}
    for row in ws.iter_rows(min_row=5, max_col=18, values_only=True):
        if isinstance(row[0], str) and re.fullmatch(r"S\d{3}", row[0]):
            vals = [str(v).strip().lower() if v is not None else "" for v in row[1:17]]
            rows[row[0]] = {"is_energy": vals[0], **{f"{g}:{l}": v for (g, l), v in zip(LABELS, vals[1:])}}
    return pd.DataFrame.from_dict(rows, orient="index")


def read_key():
    k = pd.read_excel(KEY_FILE).set_index("code")
    split = lambda s: set(x.strip() for x in str(s).split(",")) if isinstance(s, str) and s.strip() else set()
    out = pd.DataFrame(index=k.index)
    out["is_energy"] = k["LLM_is_energy"].map(lambda v: bool(v))
    for g, l in LABELS:
        col = {"Topic": "LLM_topics", "Risk": "LLM_risk_types", "Source": "LLM_energy_sources"}[g]
        out[f"{g}:{l}"] = k[col].map(lambda s, l=l: bool(l in split(s)))
    return out


def metrics(ref, pred):
    ref, pred = [bool(x) for x in ref], [bool(x) for x in pred]
    n = len(ref)
    if n == 0:
        return {"n": 0}
    tp = sum(a and b for a, b in zip(ref, pred)); fp = sum((not a) and b for a, b in zip(ref, pred))
    fn = sum(a and (not b) for a, b in zip(ref, pred)); tn = n - tp - fp - fn
    prec = tp / (tp + fp) if tp + fp else float("nan")
    rec = tp / (tp + fn) if tp + fn else float("nan")
    f1 = 2 * prec * rec / (prec + rec) if prec == prec and rec == rec and prec + rec else float("nan")
    return {"n": n, "Ref yes": sum(ref), "Pred yes": sum(pred), "TP": tp, "FP": fp, "FN": fn, "TN": tn,
            "Precision": prec, "Recall": rec, "F1": f1}


def compare(ref, pred, ref_name, pred_name):
    rows = []
    both = ref.index.intersection(pred.index)
    r_rel = ref.loc[both, "is_energy"]; p_rel = pred.loc[both, "is_energy"]
    ok = r_rel.notna() & p_rel.notna()
    rows.append({"Dimension": "Relevance", "Label": "energy main subject",
                 **metrics(r_rel[ok].tolist(), p_rel[ok].tolist())})
    energy = [c for c in both if ref.at[c, "is_energy"] == True and pred.at[c, "is_energy"] == True]
    for g, l in LABELS:
        rows.append({"Dimension": g, "Label": l,
                     **metrics(ref.loc[energy, f"{g}:{l}"].tolist(), pred.loc[energy, f"{g}:{l}"].tolist())})
    df = pd.DataFrame(rows)
    for g in ["Topic", "Risk", "Source"]:
        sub = df[df["Dimension"] == g]
        avg = {"Dimension": g, "Label": "AVERAGE (macro)", "n": int(sub["n"].iloc[0])}
        for m in ["Precision", "Recall", "F1"]:
            avg[m] = sub[m].mean()
        df = pd.concat([df, pd.DataFrame([avg])], ignore_index=True)
    order = {"Relevance": 0, "Topic": 1, "Risk": 2, "Source": 3}
    df = df.sort_values(by=["Dimension", "Label"], key=lambda s: s.map(order) if s.name == "Dimension" else s.str.startswith("AVERAGE"), kind="stable")
    df.insert(0, "Reference", ref_name); df.insert(1, "Compared", pred_name)
    return df.round(3)


def to_bool(df):
    out = df.copy()
    for c in out.columns:
        out[c] = out[c].map({"yes": True, "no": False, "": None})
    return out


def check(name, raw):
    issues = []
    for code, r in raw.iterrows():
        bad = [c for c, v in r.items() if v not in ("yes", "no", "")]
        if bad: issues.append(f"{name} {code}: invalid values in {bad}")
        if r["is_energy"] == "":
            issues.append(f"{name} {code}: Q1 not answered")
        elif r["is_energy"] == "yes":
            missing = [c for c, v in r.items() if c != "is_energy" and v == ""]
            if missing: issues.append(f"{name} {code}: {len(missing)} label(s) left empty, counted as 'no'")
    return issues


def main():
    llm = read_key()
    coders, issues = {}, []
    for name, path in CODER_FILES.items():
        raw = read_coder(path)
        issues += check(name, raw)
        c = to_bool(raw)
        lab = [x for x in c.columns if x != "is_energy"]
        c.loc[c["is_energy"] == True, lab] = c.loc[c["is_energy"] == True, lab].fillna(False)
        coders[name] = c

    results = [compare(c, llm, name, "LLM") for name, c in coders.items()]
    names = list(coders)
    if len(names) == 2:
        results.append(compare(coders[names[0]], coders[names[1]], names[0], names[1]))
    out = pd.concat(results, ignore_index=True)

    with pd.ExcelWriter(OUTPUT_FILE) as xw:
        out.to_excel(xw, sheet_name="Results", index=False)
        pd.DataFrame({"Data checks": issues or ["No problems found"]}).to_excel(xw, sheet_name="Data checks", index=False)

    summary = out[out["Label"].isin(["AVERAGE (macro)", "energy main subject"])]
    print(summary[["Reference", "Compared", "Dimension", "Label", "n", "Precision", "Recall", "F1"]].to_string(index=False))
    print(f"\n{len(issues)} data check message(s). Full results in {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
